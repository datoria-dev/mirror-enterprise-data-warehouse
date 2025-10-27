"""
Generate Excel file with ERD documentation from Snowflake metadata
"""

import json
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Read the latest metadata extraction results
results_dir = Path('QUERY_RESULTS')
latest_file = sorted(results_dir.glob('results_EXTRACT_COMPLETE_ERD_METADATA_*.json'))[-1]

print(f"Reading metadata from: {latest_file}")

with open(latest_file, 'r', encoding='utf-8') as f:
    data = json.load(f)

# Create Excel workbook
wb = Workbook()
wb.remove(wb.active)  # Remove default sheet

# Styles
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)
subheader_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
subheader_font = Font(bold=True, size=10)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

def style_header(ws, row=1, cols=None):
    """Apply header styling to a row"""
    if cols is None:
        cols = ws.max_column
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

def auto_size_columns(ws):
    """Auto-size columns based on content"""
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

# ============================================================================
# SHEET 1: Summary
# ============================================================================
ws_summary = wb.create_sheet("Architecture Summary")

summary_data = [
    ["SECURITY_ANALYTICS Data Warehouse - Architecture Summary"],
    ["Generated:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
    ["Environment:", "DEV (Development)"],
    ["Cloud Platform:", "Microsoft Azure Snowflake"],
    ["Account:", "mw76572.east-us-2.azure"],
    [],
    ["Layer", "Tables", "Views", "Procedures", "Tasks", "Total Objects"],
    ["DEV_LANDING", 141, 11, 0, 0, 152],
    ["DEV_TRANSFORMATION", 117, "50+", 15, 2, "184+"],
    ["DEV_REPORTING", 18, 148, 10, 0, 176],
    ["TOTAL", 276, "209+", 25, 2, "512+"],
    [],
    ["Key Metrics"],
    ["Primary Keys", 72],
    ["Foreign Keys", 17],
    ["Documented Tables", 45],
    ["Data Size (approx)", "~28 GB"],
    ["Dimension Tables", 32],
    ["Fact Tables", 23],
]

for row_idx, row_data in enumerate(summary_data, 1):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws_summary.cell(row=row_idx, column=col_idx, value=value)
        if row_idx == 1:
            cell.font = Font(bold=True, size=14, color="366092")
        elif row_idx == 7 or row_idx == 13:
            cell.font = subheader_font
            cell.fill = subheader_fill

style_header(ws_summary, 7, 6)
auto_size_columns(ws_summary)

# ============================================================================
# SHEET 2: Layer 1 - Landing Tables
# ============================================================================
ws_landing = wb.create_sheet("Layer 1 - Landing")

# Extract landing tables from results
landing_tables = []
for stmt in data['statements']:
    if stmt.get('sql', '').startswith("SELECT\n    'LANDING' as LAYER,\n    'TABLE'"):
        for row in stmt.get('rows', []):
            landing_tables.append([
                row.get('OBJECT_NAME'),
                row.get('ROW_COUNT'),
                row.get('SIZE_MB'),
                row.get('COMMENT', '')
            ])
        break

ws_landing.append(["Landing Zone Tables (DEV_LANDING.SECURITY_ANALYTICS)"])
ws_landing.append([])
ws_landing.append(["Table Name", "Row Count", "Size (MB)", "Description"])

for table in landing_tables:
    ws_landing.append(table)

ws_landing.cell(1, 1).font = Font(bold=True, size=14, color="366092")
style_header(ws_landing, 3, 4)
auto_size_columns(ws_landing)

# ============================================================================
# SHEET 3: Layer 2 - Dimensions
# ============================================================================
ws_dimensions = wb.create_sheet("Layer 2 - Dimensions")

dimensions = []
for stmt in data['statements']:
    if "'DIMENSION' as OBJECT_TYPE" in stmt.get('sql', ''):
        for row in stmt.get('rows', []):
            dimensions.append([
                row.get('OBJECT_NAME'),
                row.get('ROW_COUNT'),
                row.get('SIZE_MB'),
                row.get('COMMENT', '')
            ])
        break

ws_dimensions.append(["Dimension Tables (DEV_TRANSFORMATION.SECURITY_ANALYTICS)"])
ws_dimensions.append([])
ws_dimensions.append(["Dimension Name", "Row Count", "Size (MB)", "Description"])

for dim in dimensions:
    ws_dimensions.append(dim)

ws_dimensions.cell(1, 1).font = Font(bold=True, size=14, color="366092")
style_header(ws_dimensions, 3, 4)
auto_size_columns(ws_dimensions)

# ============================================================================
# SHEET 4: Layer 2 - Facts
# ============================================================================
ws_facts = wb.create_sheet("Layer 2 - Facts")

facts = []
for stmt in data['statements']:
    if "'FACT' as OBJECT_TYPE" in stmt.get('sql', ''):
        for row in stmt.get('rows', []):
            facts.append([
                row.get('OBJECT_NAME'),
                row.get('ROW_COUNT'),
                row.get('SIZE_MB'),
                row.get('COMMENT', '')
            ])
        break

ws_facts.append(["Fact Tables (DEV_TRANSFORMATION.SECURITY_ANALYTICS)"])
ws_facts.append([])
ws_facts.append(["Fact Name", "Row Count", "Size (MB)", "Description"])

for fact in facts:
    ws_facts.append(fact)

ws_facts.cell(1, 1).font = Font(bold=True, size=14, color="366092")
style_header(ws_facts, 3, 4)
auto_size_columns(ws_facts)

# ============================================================================
# SHEET 5: Key Dimensions - Detailed Structure
# ============================================================================
ws_key_dims = wb.create_sheet("Key Dimensions Structure")

key_dimensions = ['DIM_OPCO', 'DIM_HOST', 'DIM_DATES', 'DIM_USER']
current_row = 1

for dim_name in key_dimensions:
    # Find the dimension structure in results
    for stmt in data['statements']:
        if f"'{dim_name}' as TABLE_NAME" in stmt.get('sql', ''):
            ws_key_dims.cell(current_row, 1, dim_name).font = subheader_font
            ws_key_dims.cell(current_row, 1).fill = subheader_fill
            current_row += 1

            ws_key_dims.append(["Column Name", "Data Type", "Nullable"])
            style_header(ws_key_dims, current_row, 3)
            current_row += 1

            for row in stmt.get('rows', []):
                ws_key_dims.append([
                    row.get('COLUMN_NAME'),
                    row.get('DATA_TYPE'),
                    row.get('IS_NULLABLE')
                ])
                current_row += 1

            current_row += 1
            break

auto_size_columns(ws_key_dims)

# ============================================================================
# SHEET 6: Key Facts - Detailed Structure
# ============================================================================
ws_key_facts = wb.create_sheet("Key Facts Structure")

key_facts = ['FACT_EDR', 'FACT_QUALYS']
current_row = 1

for fact_name in key_facts:
    for stmt in data['statements']:
        if f"'{fact_name}' as TABLE_NAME" in stmt.get('sql', ''):
            ws_key_facts.cell(current_row, 1, fact_name).font = subheader_font
            ws_key_facts.cell(current_row, 1).fill = subheader_fill
            current_row += 1

            ws_key_facts.append(["Column Name", "Data Type", "Nullable"])
            style_header(ws_key_facts, current_row, 3)
            current_row += 1

            for row in stmt.get('rows', []):
                ws_key_facts.append([
                    row.get('COLUMN_NAME'),
                    row.get('DATA_TYPE'),
                    row.get('IS_NULLABLE')
                ])
                current_row += 1

            current_row += 1
            break

auto_size_columns(ws_key_facts)

# ============================================================================
# SHEET 7: Layer 3 - Reporting Views
# ============================================================================
ws_reporting = wb.create_sheet("Layer 3 - Reporting")

reporting_views = []
for stmt in data['statements']:
    if "'REPORTING' as LAYER" in stmt.get('sql', '') and "'VIEW' as OBJECT_TYPE" in stmt.get('sql', ''):
        for row in stmt.get('rows', []):
            reporting_views.append([
                row.get('OBJECT_NAME'),
                row.get('COMMENT', '')
            ])

ws_reporting.append(["Reporting Views (DEV_REPORTING.SECURITY_ANALYTICS)"])
ws_reporting.append([])
ws_reporting.append(["View Name", "Description"])

for view in reporting_views[:50]:  # First 50
    ws_reporting.append(view)

ws_reporting.cell(1, 1).font = Font(bold=True, size=14, color="366092")
style_header(ws_reporting, 3, 2)
auto_size_columns(ws_reporting)

# ============================================================================
# SHEET 8: Stored Procedures
# ============================================================================
ws_procedures = wb.create_sheet("Stored Procedures")

procedures = []
for stmt in data['statements']:
    if "'PROCEDURE' as OBJECT_TYPE" in stmt.get('sql', ''):
        for row in stmt.get('rows', []):
            procedures.append([
                row.get('LAYER'),
                row.get('OBJECT_NAME'),
                row.get('ARGUMENT_SIGNATURE'),
                row.get('PROCEDURE_LANGUAGE')
            ])

ws_procedures.append(["Stored Procedures - All Layers"])
ws_procedures.append([])
ws_procedures.append(["Layer", "Procedure Name", "Arguments", "Language"])

for proc in procedures:
    ws_procedures.append(proc)

ws_procedures.cell(1, 1).font = Font(bold=True, size=14, color="366092")
style_header(ws_procedures, 3, 4)
auto_size_columns(ws_procedures)

# ============================================================================
# SHEET 9: Relationships (PKs & FKs)
# ============================================================================
ws_relationships = wb.create_sheet("Relationships")

pks = []
fks = []

for stmt in data['statements']:
    sql = stmt.get('sql', '')
    if 'PRIMARY KEY' in sql and 'TABLE_NAME,' in sql:
        for row in stmt.get('rows', []):
            pks.append([row.get('TABLE_NAME'), row.get('PK_NAME')])
    elif 'FOREIGN KEY' in sql and 'TABLE_NAME,' in sql:
        for row in stmt.get('rows', []):
            fks.append([row.get('TABLE_NAME'), row.get('FK_NAME')])

ws_relationships.append(["Primary Keys"])
ws_relationships.append(["Table Name", "Constraint Name"])
style_header(ws_relationships, 2, 2)

for pk in pks:
    ws_relationships.append(pk)

current_row = len(pks) + 4
ws_relationships.cell(current_row, 1, "Foreign Keys").font = subheader_font
ws_relationships.cell(current_row, 1).fill = subheader_fill
current_row += 1
ws_relationships.append(["Table Name", "Constraint Name"])
style_header(ws_relationships, current_row, 2)

for fk in fks:
    ws_relationships.append(fk)

auto_size_columns(ws_relationships)

# ============================================================================
# SHEET 10: Data Lineage
# ============================================================================
ws_lineage = wb.create_sheet("Data Lineage")

# Note: Lineage data would be in EXTRACT_ERD_METADATA.sql results
ws_lineage.append(["Data Lineage Catalog"])
ws_lineage.append([])
ws_lineage.append(["Source DB", "Source Schema", "Source Table", "Target DB", "Target Schema", "Target Table", "Transformation", "Frequency"])

lineage_data = [
    ["DEV_LANDING", "SECURITY_ANALYTICS", "L_CROWDSTRIKE_HOSTS", "DEV_TRANSFORMATION", "SECURITY_ANALYTICS", "DIM_HOST", "Merge CrowdStrike endpoints", "Daily"],
    ["DEV_LANDING", "SECURITY_ANALYTICS", "L_EDR_THREATS_REALTIME", "DEV_TRANSFORMATION", "SECURITY_ANALYTICS", "FACT_EDR", "Aggregate threat counts", "Real-time"],
    ["DEV_LANDING", "SECURITY_ANALYTICS", "L_PAM_USERS", "DEV_TRANSFORMATION", "SECURITY_ANALYTICS", "DIM_USER", "SCD Type 2 merge", "Daily"],
    ["DEV_TRANSFORMATION", "SECURITY_ANALYTICS", "DIM_HOST", "DEV_REPORTING", "SECURITY_ANALYTICS", "VW_POWERBI_EXECUTIVE_DASHBOARD", "Join with facts", "On-demand"],
    ["DEV_TRANSFORMATION", "SECURITY_ANALYTICS", "FACT_EDR", "DEV_REPORTING", "SECURITY_ANALYTICS", "VW_EXECUTIVE_KPI_DASHBOARD", "Calculate KPIs", "On-demand"],
]

for row in lineage_data:
    ws_lineage.append(row)

ws_lineage.cell(1, 1).font = Font(bold=True, size=14, color="366092")
style_header(ws_lineage, 3, 8)
auto_size_columns(ws_lineage)

# Save workbook
output_file = f'ITSECKPI_ERD_Documentation_{datetime.now().strftime("%Y%m%d")}.xlsx'
wb.save(output_file)

print(f"\n✓ Excel file created: {output_file}")
print(f"\nSheets created:")
print(f"  1. Architecture Summary")
print(f"  2. Layer 1 - Landing ({len(landing_tables)} tables)")
print(f"  3. Layer 2 - Dimensions ({len(dimensions)} tables)")
print(f"  4. Layer 2 - Facts ({len(facts)} tables)")
print(f"  5. Key Dimensions Structure")
print(f"  6. Key Facts Structure")
print(f"  7. Layer 3 - Reporting ({len(reporting_views)} views)")
print(f"  8. Stored Procedures ({len(procedures)} procedures)")
print(f"  9. Relationships ({len(pks)} PKs, {len(fks)} FKs)")
print(f"  10. Data Lineage")
