#!/usr/bin/env python3
"""
Extract text content from PowerPoint presentations for analysis
"""

from pptx import Presentation
import json
import sys

def extract_pptx_content(pptx_path):
    """Extract all text content from a PPTX file"""
    try:
        prs = Presentation(pptx_path)

        presentation_data = {
            'file_name': pptx_path.split('\\')[-1],
            'total_slides': len(prs.slides),
            'slides': []
        }

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_data = {
                'slide_number': slide_num,
                'title': '',
                'content': []
            }

            # Extract text from all shapes
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        # Try to identify if it's a title
                        if shape.is_placeholder and shape.placeholder_format.type == 1:  # Title placeholder
                            slide_data['title'] = text
                        else:
                            slide_data['content'].append(text)

            presentation_data['slides'].append(slide_data)

        return presentation_data

    except Exception as e:
        return {'error': str(e)}

def main():
    pptx_files = [
        r'C:\Users\fonat\OneDrive\Documents\GenericCorp\Snowflake_ITSECKPI_Project\GIS Offsite Event - Snowflake.pptx',
        r'C:\Users\fonat\OneDrive\Documents\GenericCorp\Snowflake_ITSECKPI_Project\GIS-Data-Platform.pptx'
    ]

    for pptx_file in pptx_files:
        print(f"\n{'='*80}")
        print(f"Extracting: {pptx_file.split('\\')[-1]}")
        print('='*80)

        data = extract_pptx_content(pptx_file)

        if 'error' in data:
            print(f"ERROR: {data['error']}")
            continue

        print(f"\nTotal Slides: {data['total_slides']}\n")

        for slide in data['slides']:
            print(f"\n--- SLIDE {slide['slide_number']} ---")
            if slide['title']:
                print(f"TITLE: {slide['title']}")
            if slide['content']:
                print("CONTENT:")
                for content in slide['content']:
                    print(f"  • {content}")

        # Save to JSON for detailed analysis
        output_file = pptx_file.replace('.pptx', '_extracted.json')
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f"\n\nDetailed content saved to: {output_file}")

if __name__ == '__main__':
    main()
